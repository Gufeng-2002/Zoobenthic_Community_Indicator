#!/usr/bin/env python3
"""
Convert Excel tables from analysis results into LaTeX format.

Walks through the results directory tree, finds all tables/ subfolders,
reads each .xlsx file, and writes a corresponding .tex file to
a latex_tables/ folder alongside the tables/ folder.

Usage:
    python convert_tables_to_latex.py
"""

import os
import re
import openpyxl

RESULTS_DIR = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------------------
# Caption mapping:  stage_path/base_name  →  caption text
# For cluster-specific tables, use {cluster_num} placeholder.
# ---------------------------------------------------------------------------
CAPTION_MAP = {
    # 01 Pollution Assessment
    "01_pollution_assessment/pc_loadings":
        "Principal component loadings for pollution-indicator variables",
    "01_pollution_assessment/site_scores":
        "Site pollution scores across the first five principal components",

    # 02 Taxa Assemblage
    "02_taxa_assemblage/anova_env":
        "One-way ANOVA results for environmental variables across reference-site clusters",
    "02_taxa_assemblage/anova_taxa":
        "One-way ANOVA results for taxa relative abundance across reference-site clusters",
    "02_taxa_assemblage/reference_taxa_clusters":
        "Octave-transformed taxa abundance at reference sites by cluster membership",

    # 03 LDA Classification
    "03_LDA_Classification/lda_axes_summary":
        "Summary of Linear Discriminant Analysis (LDA) axes",
    "03_LDA_Classification/lda_classification_report":
        "LDA classification report for cluster assignment on the full dataset",
    "03_LDA_Classification/lda_confusion_matrix":
        "LDA confusion matrix for cluster prediction on the full dataset",
    "03_LDA_Classification/lda_env_significance":
        "Environmental variable significance across LDA-defined clusters",
    "03_LDA_Classification/mccv_classification_report":
        "Monte Carlo cross-validation classification report (1,000 iterations, 20\\% test size)",
    "03_LDA_Classification/mccv_confusion_matrix":
        "Monte Carlo cross-validation confusion matrix (cumulative over 1,000 iterations)",

    # 04 Bray-Curtis / NMDS
    "04_bray_curtis_NMDS/nmds_summary":
        "Non-metric multidimensional scaling (NMDS) ordination summary by cluster",
    "04_bray_curtis_NMDS/zci_summary":
        "Zooplankton Community Index (ZCI) summary statistics and correlation "
        "with pollution scores by cluster",

    # 05 Piecewise Quantile Regression
    "05_piecewise_qr/qr_coefficients_cluster":
        "Piecewise quantile regression coefficients for Cluster~{cluster_num} "
        "across quantile levels ($\\tau = 0.10$--$0.90$)",
    "05_piecewise_qr/sensitivity_cluster":
        "Subsample sensitivity analysis of piecewise quantile regression "
        "for Cluster~{cluster_num}",

    # RDA
    "RDA_analysis/rda_axes_summary":
        "Redundancy Analysis (RDA) axes summary with eigenvalues and significance tests",
    "RDA_analysis/rda_terms_summary":
        "RDA marginal (Type~III) tests for environmental predictor variables",
}

# Detroit River variants share the same base captions with a suffix
DR_SUFFIX = " (Detroit River)"


# ──────────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────────

def escape_latex(text):
    """Escape characters that are special in LaTeX text mode."""
    if text is None:
        return ""
    text = str(text)
    # Characters to escape (order matters: backslash first)
    text = text.replace('\\', r'\textbackslash{}')
    for ch, repl in [('&', r'\&'), ('%', r'\%'), ('$', r'\$'),
                      ('#', r'\#'), ('_', r'\_'), ('{', r'\{'),
                      ('}', r'\}'), ('~', r'\textasciitilde{}'),
                      ('^', r'\textasciicircum{}')]:
        text = text.replace(ch, repl)
    # Replace ± with LaTeX math
    text = text.replace('±', '$\\pm$')
    # Replace degree symbol
    text = text.replace('°', '$^{\\circ}$')
    # oC → $^{\\circ}$C  (common in the data)
    text = re.sub(r'\(oC\)', r'($^{\\circ}$C)', text)
    return text


def format_number(val, decimals=4):
    """Format a numeric value for LaTeX, choosing sensible precision."""
    if isinstance(val, bool):
        return str(val)
    if isinstance(val, int):
        return f"{val:,}".replace(",", "{,}")  # thousands separator
    if isinstance(val, float):
        abs_v = abs(val)
        if abs_v == 0:
            return "0"
        if abs_v < 1e-6:
            # Scientific notation
            return f"${val:.2e}$".replace("e-0", r" \times 10^{-").replace("e+0", r" \times 10^{") \
                   .replace("e-", r" \times 10^{-").replace("e+", r" \times 10^{") + "}$" \
                   if False else f"{val:.2e}"
        if abs_v < 0.001:
            return f"{val:.6f}"
        if abs_v < 1:
            return f"{val:.{decimals}f}"
        if abs_v < 100:
            return f"{val:.{min(decimals, 2)}f}"
        if abs_v < 10000:
            return f"{val:.{min(decimals, 1)}f}"
        return f"{val:.0f}"
    return str(val)


def format_cell(val):
    """Format a single cell value for LaTeX."""
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return ""
    if isinstance(val, (int, float)):
        return format_number(val)
    # String value — escape LaTeX specials
    return escape_latex(str(val))


def is_index_column(ws):
    """Return True if column A is a 0-based or 1-based sequential integer index.

    Heuristic: column A is an index when its header cell (A1) is empty/None
    AND the remaining values form a consecutive integer sequence.
    If A1 carries a meaningful header (e.g. 'Cluster', 'StationID'),
    the column is real data, not an index.
    """
    header_val = ws.cell(row=1, column=1).value
    if header_val is not None and str(header_val).strip() != "":
        return False          # meaningful header → not a throw-away index

    vals = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        v = row[0]
        if v is not None:
            try:
                vals.append(int(v))
            except (ValueError, TypeError):
                return False
    if len(vals) < 2:
        return False
    diffs = [vals[i+1] - vals[i] for i in range(len(vals)-1)]
    return all(d == 1 for d in diffs)


# ──────────────────────────────────────────────────────────────────────────────
# Caption / label helpers
# ──────────────────────────────────────────────────────────────────────────────

def _caption_key(stage_path, filename):
    """Build the lookup key into CAPTION_MAP."""
    name = os.path.splitext(filename)[0]

    # Strip trailing _N for cluster-specific tables
    cluster_match = re.search(r'_(\d+)$', name)
    base = re.sub(r'_\d+$', '', name) if cluster_match else name

    # Remove DR_results/ prefix for lookup
    clean_stage = stage_path
    if clean_stage.startswith("DR_results/"):
        clean_stage = clean_stage[len("DR_results/"):]

    return clean_stage + "/" + base, cluster_match, stage_path.startswith("DR_results")


def get_caption(stage_path, filename):
    key, cluster_match, is_dr = _caption_key(stage_path, filename)
    caption = CAPTION_MAP.get(key)
    if caption is None:
        # Fallback: prettify filename
        caption = os.path.splitext(filename)[0].replace('_', ' ').title()
    if cluster_match:
        caption = caption.replace("{cluster_num}", cluster_match.group(1))
    if is_dr:
        caption += DR_SUFFIX
    return caption


def get_label(stage_path, filename):
    name = os.path.splitext(filename)[0]
    tag = stage_path.replace("/", "_").replace("DR_results_", "dr_")
    return f"tab:{tag}_{name}"


# ──────────────────────────────────────────────────────────────────────────────
# Special: pc_loadings table
# ──────────────────────────────────────────────────────────────────────────────

# Chemical variable groupings (by dominant PC loading).
# A separator is inserted after each group.
PC_LOADINGS_GROUPS = [
    ["Co", "Al", "Ni", "Mn", "Fe", "Cr", "Cu"],
    ["Hg", "Pb", "Zn", "total PCB", "Cd"],
    ["OCS", "p,p'-DDE"],
    ["As"],
    ["Ca"],
]

# Summary rows that should be displayed with two-line labels
PC_LOADINGS_SUMMARY = {
    "Explained Variance":    r"\makecell[l]{Explained\\Variance}",
    "Proportion of Variance": r"\makecell[l]{Proportion of\\total variance}",
    "Cumulative Proportion":  r"\makecell[l]{Cum.\\Proportion}",
}


def _fmt4(val):
    """Format a numeric value to exactly 4 decimal places."""
    if val is None:
        return ""
    if isinstance(val, str):
        return escape_latex(val)
    return f"{val:.4f}"


def convert_pc_loadings(xlsx_path, stage_path, filename):
    """Specialised converter for pc_loadings tables."""
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    if not all_rows:
        return None

    # Build {variable_name: [pc1..pc5]} mapping
    header = all_rows[0]  # e.g. [None, 'PC1', 'PC2', ...]
    pc_headers = [str(h) for h in header[1:] if h is not None]
    n_pc = len(pc_headers)

    row_dict = {}
    for row in all_rows[1:]:
        name = row[0]
        if name is None:
            continue
        name = str(name).strip()
        row_dict[name] = row[1:1 + n_pc]

    caption = get_caption(stage_path, filename)
    label  = get_label(stage_path, filename)

    lines = []
    lines.append(r"% Requires: \usepackage{booktabs, makecell}")
    lines.append(r"\begin{table}[htbp]")
    lines.append(r"  \centering")
    lines.append(f"  \\caption{{{caption}}}")
    lines.append(f"  \\label{{{label}}}")
    col_spec = "l" + "r" * n_pc
    lines.append(f"  \\begin{{tabular}}{{{col_spec}}}")
    lines.append(r"    \toprule")
    hdr = "    \\textbf{Chemical variables}"
    for pc in pc_headers:
        hdr += f" & \\textbf{{{pc}}}"
    hdr += r" \\"
    lines.append(hdr)
    lines.append(r"    \midrule")

    # Chemical variable groups
    for gi, group in enumerate(PC_LOADINGS_GROUPS):
        for name in group:
            vals = row_dict.get(name)
            if vals is None:
                continue
            label_cell = f"    \\textbf{{{escape_latex(name)}}}"
            cells = [label_cell] + [_fmt4(v) for v in vals]
            lines.append(" & ".join(cells) + r" \\")
        # Separator after each group (but not the very last chemical group)
        if gi < len(PC_LOADINGS_GROUPS) - 1:
            lines.append(r"    \addlinespace")

    # Separator before summary rows
    lines.append(r"    \midrule")

    # Summary rows with two-line labels
    for name, latex_label in PC_LOADINGS_SUMMARY.items():
        vals = row_dict.get(name)
        if vals is None:
            continue
        cells = [f"    {latex_label}"] + [_fmt4(v) for v in vals]
        lines.append(" & ".join(cells) + r" \\")

    lines.append(r"    \bottomrule")
    lines.append(r"  \end{tabular}")
    lines.append(r"\end{table}")

    return "\n".join(lines) + "\n"


# ──────────────────────────────────────────────────────────────────────────────
# Core conversion
# ──────────────────────────────────────────────────────────────────────────────

def _classify_rows(all_rows, start_col, num_cols):
    """
    Separate data rows from note/empty rows.

    Returns
    -------
    data_rows : list[list]   – rows to render inside the tabular body
    notes     : list[str]    – long-string rows to render as table notes
    """
    data_rows = []
    notes = []

    for i, row in enumerate(all_rows):
        if i == 0:
            # header — always keep
            data_rows.append(row)
            continue

        relevant = row[start_col:]
        non_none = [(j, v) for j, v in enumerate(relevant) if v is not None
                    and str(v).strip() != ""]

        # Completely empty row → skip
        if len(non_none) == 0:
            continue

        # Single cell with a long text → treat as footnote / note
        if len(non_none) == 1:
            _, val = non_none[0]
            text = str(val).strip()
            if len(text) > 45:
                notes.append(text)
                continue

        data_rows.append(row)

    return data_rows, notes


def _col_alignment(header, num_cols):
    """Choose column alignments: first col left, numbers right-ish."""
    return "l" + "r" * (num_cols - 1) if num_cols > 1 else "l"


def _needs_longtable(n_rows):
    return n_rows > 40


def convert_one(xlsx_path, stage_path, filename):
    """Return LaTeX source string for one Excel table."""

    # Special handling for pc_loadings
    if os.path.splitext(filename)[0] == "pc_loadings":
        return convert_pc_loadings(xlsx_path, stage_path, filename)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    if not all_rows:
        return None

    skip_first = is_index_column(ws)
    start_col = 1 if skip_first else 0
    num_cols = len(all_rows[0]) - start_col

    data_rows, notes = _classify_rows(all_rows, start_col, num_cols)
    if len(data_rows) < 2:          # header only
        return None

    header = data_rows[0][start_col:]
    header_cells = [escape_latex(str(h)) if h else "" for h in header]
    body = data_rows[1:]

    caption = get_caption(stage_path, filename)
    label  = get_label(stage_path, filename)
    align  = _col_alignment(header_cells, num_cols)

    use_long = _needs_longtable(len(body))
    use_resize = num_cols > 10
    use_small = 7 <= num_cols <= 10

    lines = []

    # ── preamble ──────────────────────────────────────────────────────────
    if use_long:
        lines.append(r"% Requires: \usepackage{longtable, booktabs}")
        if use_small:
            lines.append(r"{\small")
        if use_resize:
            lines.append(r"{\footnotesize")

        lines.append(f"\\begin{{longtable}}{{{align}}}")
        lines.append(f"  \\caption{{{caption}}}")
        lines.append(f"  \\label{{{label}}} \\\\")
        lines.append(r"  \toprule")
        lines.append("  " + " & ".join(header_cells) + r" \\")
        lines.append(r"  \midrule")
        lines.append(r"  \endfirsthead")
        # Continuation header
        lines.append(f"  \\multicolumn{{{num_cols}}}{{l}}"
                      f"{{\\tablename\\ \\thetable\\ -- continued}} \\\\")
        lines.append(r"  \toprule")
        lines.append("  " + " & ".join(header_cells) + r" \\")
        lines.append(r"  \midrule")
        lines.append(r"  \endhead")
        lines.append(f"  \\midrule \\multicolumn{{{num_cols}}}{{r}}"
                      r"{{Continued on next page}} \\")
        lines.append(r"  \endfoot")
        lines.append(r"  \bottomrule")
        # Notes in the last footer
        if notes:
            for note in notes:
                lines.append(
                    f"  \\multicolumn{{{num_cols}}}{{l}}"
                    f"{{\\footnotesize {escape_latex(note)}}} \\\\"
                )
        lines.append(r"  \endlastfoot")
    else:
        lines.append(r"\begin{table}[htbp]")
        lines.append(r"  \centering")
        lines.append(f"  \\caption{{{caption}}}")
        lines.append(f"  \\label{{{label}}}")
        if use_resize:
            lines.append(r"  \resizebox{\textwidth}{!}{%")
        elif use_small:
            lines.append(r"  \small")
        lines.append(f"  \\begin{{tabular}}{{{align}}}")
        lines.append(r"    \toprule")
        lines.append("    " + " & ".join(header_cells) + r" \\")
        lines.append(r"    \midrule")

    # ── body ──────────────────────────────────────────────────────────────
    indent = "  " if use_long else "    "
    for row in body:
        cells = row[start_col:]
        formatted = [format_cell(v) for v in cells]
        lines.append(indent + " & ".join(formatted) + r" \\")

    # ── closing ───────────────────────────────────────────────────────────
    if use_long:
        lines.append(r"\end{longtable}")
        if use_small or use_resize:
            lines.append("}")
    else:
        lines.append(r"    \bottomrule")
        lines.append(r"  \end{tabular}")
        if use_resize:
            lines.append(r"  }")
        # Table notes
        if notes:
            lines.append(r"  \\[4pt]")
            for note in notes:
                lines.append(
                    f"  \\parbox{{\\textwidth}}"
                    f"{{\\footnotesize {escape_latex(note)}}}"
                )
        lines.append(r"\end{table}")

    return "\n".join(lines) + "\n"


# ──────────────────────────────────────────────────────────────────────────────
# Directory walker
# ──────────────────────────────────────────────────────────────────────────────

def walk_and_convert(results_dir):
    """Find all tables/ folders, convert .xlsx → .tex in latex_tables/."""
    converted = 0
    for dirpath, dirnames, filenames in os.walk(results_dir):
        # Only process directories named "tables"
        if os.path.basename(dirpath) != "tables":
            continue

        xlsx_files = sorted(f for f in filenames if f.endswith(".xlsx"))
        if not xlsx_files:
            continue

        # Create latex_tables/ alongside tables/
        parent = os.path.dirname(dirpath)
        latex_dir = os.path.join(parent, "latex_tables")
        os.makedirs(latex_dir, exist_ok=True)

        # Compute stage_path relative to results_dir
        stage_path = os.path.relpath(parent, results_dir)

        for fname in xlsx_files:
            xlsx_path = os.path.join(dirpath, fname)
            tex_name = os.path.splitext(fname)[0] + ".tex"
            tex_path = os.path.join(latex_dir, tex_name)

            tex_src = convert_one(xlsx_path, stage_path, fname)
            if tex_src is None:
                print(f"  [skip] {stage_path}/tables/{fname}  (empty)")
                continue

            with open(tex_path, "w", encoding="utf-8") as fh:
                fh.write(tex_src)
            converted += 1
            print(f"  [ok]   {stage_path}/latex_tables/{tex_name}")

    return converted


# ──────────────────────────────────────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print(f"Results directory: {RESULTS_DIR}\n")
    n = walk_and_convert(RESULTS_DIR)
    print(f"\nDone — {n} table(s) converted.")
